import os
import sys
import time
import cProfile
import pstats
import random
import openpyxl
import sqlite3
from datetime import datetime, timedelta

# Add parent directory to sys.path to import local modules
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import database
import config

def generate_fake_excel(filename, num_rows=5000):
    print(f"Generating fake Excel with {num_rows} rows: {filename}...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Name", "Spec", "Price"]) # Header
    
    for i in range(1, num_rows + 1):
        ws.append([
            i, 
            f"Medicine Test {i} {random.randint(1000, 9999)}", 
            random.choice(["Viên", "Vỉ", "Chai", "Ống"]), 
            random.randint(1000, 500000)
        ])
    
    wb.save(filename)
    print("Done generating Excel.")

def test_excel_import(filename):
    print("\n--- Testing Excel Import ---")
    
    # We'll simulate the task() logic from ui_medicine_pyside.py
    def import_task():
        wb = openpyxl.load_workbook(filename, data_only=True)
        s = wb.active
        
        # Fetch existing names to avoid duplicates efficiently
        existing_names = {m['name'].lower(): True for m in database.get_all_medicines_db()}
        
        to_insert = []
        for r in s.iter_rows(min_row=2, values_only=True):
            if len(r) >= 2 and r[1]:
                name = str(r[1]).strip()
                if not name or name.lower() in existing_names:
                    continue
                    
                spec = str(r[2]).strip() if len(r) > 2 and r[2] else ""
                try:
                    price_str = str(r[3]).replace(",", ".") if len(r) > 3 and r[3] is not None else "0"
                    price = float(price_str)
                except: 
                    price = 0.0
                
                to_insert.append({
                    'name': name,
                    'packing_spec': spec,
                    'price': price,
                    'stock_quantity': 0,
                    'min_stock_level': 5
                })
                existing_names[name.lower()] = True
        
        if to_insert:
            print(f"Bulk inserting {len(to_insert)} medicines...")
            database.insert_medicines_bulk(to_insert)
            return len(to_insert)
        return 0

    profiler = cProfile.Profile()
    profiler.enable()
    start_time = time.time()
    
    count = import_task()
    
    end_time = time.time()
    profiler.disable()
    
    print(f"Imported {count} medicines in {end_time - start_time:.4f} seconds.")
    
    stats = pstats.Stats(profiler).sort_stats('cumulative')
    stats.print_stats(20) # Show top 20 lines

def test_patient_history_performance(num_records=500):
    print("\n--- Testing Patient History Performance ---")
    
    # 1. Create a test patient
    patient_id = database.add_patient_db(
        "Test Patient Performance", "1990-01-01", "Nam", "Hanoi", "0987654321", "70kg", "None"
    )
    
    if not patient_id:
        print("Failed to create test patient.")
        return

    print(f"Created test patient ID: {patient_id}")
    
    # 2. Create some medicines first to link to prescriptions
    medicines = database.get_all_medicines_db()
    if not medicines:
        print("No medicines found. Adding some...")
        for i in range(5):
            database.add_medicine_db(f"Med for RX {i}", "Viên", 1000)
        medicines = database.get_all_medicines_db()
    
    med_ids = [m['id'] for m in medicines[:5]]
    
    # 3. Create 500 prescriptions bulk
    print(f"Generating {num_records} prescriptions...")
    headers = []
    base_date = datetime.now()
    
    for i in range(num_records):
        headers.append({
            'patient_id': patient_id,
            'prescription_date': (base_date - timedelta(days=i)).strftime('%Y-%m-%d %H:%M:%S'),
            'diagnosis': f"Diagnosis test number {i}",
            'total_amount': 50000,
            'notes': "Performance test note"
        })
    
    # We need to insert headers and get IDs to insert details
    # But for simplified performance test on FETCH, we can just insert them one by one or use bulk if available and then fetch.
    # We use bulk insert here.
    
    database.insert_headers_bulk(headers)
    
    # Get the header IDs we just inserted
    conn = sqlite3.connect(config.get_database_path())
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT id FROM prescriptions_header WHERE patient_id = ?", (patient_id,))
    header_ids = [row['id'] for row in c.fetchall()]
    conn.close()
    
    print(f"Inserted {len(header_ids)} headers. Inserting details...")
    
    details = []
    for h_id in header_ids:
        for m_id in med_ids:
            details.append({
                'prescription_header_id': h_id,
                'medicine_id': m_id,
                'quantity': 2,
                'unit_price': 10000
            })
    
    database.insert_details_bulk(details)
    print("Done generating records.")
    
    # 4. Measure FETCH performance
    print("Measuring get_prescriptions_by_patient_db speed...")
    
    profiler = cProfile.Profile()
    profiler.enable()
    start_time = time.time()
    
    results = database.get_prescriptions_by_patient_db(patient_id)
    
    end_time = time.time()
    profiler.disable()
    
    print(f"Fetched {len(results)} prescriptions with details in {end_time - start_time:.4f} seconds.")
    
    stats = pstats.Stats(profiler).sort_stats('cumulative')
    stats.print_stats(20)

def test_dashboard_stats():
    print("\n--- Testing Dashboard Stats ---")
    
    # Check if we can get stats by day
    # We need a valid 'YYYY-MM'
    now = datetime.now()
    ym = now.strftime('%Y-%m')
    
    print(f"Fetching stats for {ym}...")
    start_time = time.time()
    results = database.get_stats_by_day_for_month(ym)
    end_time = time.time()
    
    print(f"Got {len(results)} days of stats in {end_time - start_time:.4f} seconds.")
    
    # Check if medicine usage stats work
    print("Fetching medicine usage stats...")
    start_time = time.time()
    med_stats = database.get_medicine_usage_stats_db()
    end_time = time.time()
    print(f"Got {len(med_stats)} medicine stats in {end_time - start_time:.4f} seconds.")

def cleanup():
    print("\n--- Cleaning up ---")
    if os.path.exists("temp_medicines_5000.xlsx"):
        os.remove("temp_medicines_5000.xlsx")
    
    # Clean up DB: Delete test patient and medicines
    conn = sqlite3.connect(config.get_database_path())
    c = conn.cursor()
    c.execute("DELETE FROM patients WHERE name = 'Test Patient Performance'")
    c.execute("DELETE FROM medicines WHERE name LIKE 'Medicine Test %' OR name LIKE 'Med for RX %'")
    conn.commit()
    conn.close()
    print("Cleanup done.")

if __name__ == "__main__":
    excel_file = "temp_medicines_5000.xlsx"
    try:
        generate_fake_excel(excel_file, 5000)
        test_excel_import(excel_file)
        test_patient_history_performance(500)
        test_dashboard_stats()
    finally:
        cleanup()
